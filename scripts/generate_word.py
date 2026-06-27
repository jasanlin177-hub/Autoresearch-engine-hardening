#!/usr/bin/env python3
"""
Generate a formatted Word (.docx) report from a TaiEquityautoresearch Markdown file.
Supports Traditional Chinese, tables, headers, blockquotes, hyperlinks, bold,
and a professional cover page with live financial data.
"""
import re
import os
import sys
from datetime import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import urllib.parse

# ── Load financial data module ────────────────────────────────────────────────
_SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
if _SCRIPTS_DIR not in sys.path:
    sys.path.insert(0, _SCRIPTS_DIR)
try:
    from fetch_financials import fetch_financials, FinancialData
    HAS_FINANCIALS = True
except ImportError:
    HAS_FINANCIALS = False

# ── Helpers ───────────────────────────────────────────────────────────────────

def add_hyperlink(paragraph, text, url):
    """Add a clickable hyperlink to a paragraph."""
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    # Blue underline style
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(color)
    rPr.append(u)
    new_run.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def set_cell_bg(cell, hex_color):
    """Set table cell background color."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def apply_inline_to_para(paragraph, text, is_bold_default=False):
    """Parse inline markdown (bold, links) and add runs to paragraph."""
    # Split on **bold** and [link](url)
    token_re = re.compile(r'(\*\*(.+?)\*\*|\[([^\]]+)\]\(([^)]+)\))')
    last = 0
    for m in token_re.finditer(text):
        # Plain text before match
        if m.start() > last:
            run = paragraph.add_run(text[last:m.start()])
            if is_bold_default:
                run.bold = True
        full = m.group(0)
        if full.startswith('**'):
            run = paragraph.add_run(m.group(2))
            run.bold = True
        else:
            # hyperlink
            link_text = m.group(3)
            link_url = m.group(4)
            add_hyperlink(paragraph, link_text, link_url)
        last = m.end()
    # Remaining text
    if last < len(text):
        run = paragraph.add_run(text[last:])
        if is_bold_default:
            run.bold = True


def is_separator(line):
    return re.match(r'^-{3,}$', line.strip())

def is_table_row(line):
    return line.strip().startswith('|') and line.strip().endswith('|')

def is_table_sep(line):
    return is_table_row(line) and re.match(r'^[\|\-\:\s]+$', line.strip())

def parse_table_rows(lines):
    rows = []
    for line in lines:
        if is_table_sep(line):
            continue
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        rows.append(cells)
    return rows

def deduplicate_sections(text):
    """Keep only the last occurrence of each H1/H2 section."""
    pattern = re.compile(r'^(#{1,2} .+)$', re.MULTILINE)
    matches = list(pattern.finditer(text))
    if not matches:
        return text
    sections = []
    for i, m in enumerate(matches):
        start = m.start()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        sections.append((m.group(1), start, end))
    seen = {}
    for heading, start, end in sections:
        seen[heading] = (start, end)
    first_heading_start = sections[0][1]
    result = text[:first_heading_start]
    for heading, (start, end) in seen.items():
        result += text[start:end]
    return result


# ── Document styles ───────────────────────────────────────────────────────────

def setup_document(doc):
    """Configure default styles and margins."""
    # Page margins
    for section in doc.sections:
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # Default font (CJK compatible)
    style = doc.styles['Normal']
    style.font.name = '微軟正黑體'
    style.font.size = Pt(10)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')

    # Heading 1
    h1 = doc.styles['Heading 1']
    h1.font.name = '微軟正黑體'
    h1.font.size = Pt(15)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(0x16, 0x21, 0x3e)
    h1._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')

    # Heading 2
    h2 = doc.styles['Heading 2']
    h2.font.name = '微軟正黑體'
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(0x0f, 0x34, 0x60)
    h2._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')

    # Heading 3
    h3 = doc.styles['Heading 3']
    h3.font.name = '微軟正黑體'
    h3.font.size = Pt(11)
    h3.font.bold = True
    h3.font.color.rgb = RGBColor(0x53, 0x34, 0x83)
    h3._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')


def add_styled_para(doc, text, style_name='Normal', bold=False, italic=False,
                    color=None, left_indent=None, space_before=0, space_after=4):
    p = doc.add_paragraph(style=style_name)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after  = Pt(space_after)
    if left_indent:
        p.paragraph_format.left_indent = left_indent
    apply_inline_to_para(p, text, is_bold_default=bold)
    if color:
        for run in p.runs:
            run.font.color.rgb = color
    return p


# ── Main builder ──────────────────────────────────────────────────────────────

def build_word(md_path, out_path, ticker=None, company_name=None, market='TWO'):
    with open(md_path, encoding='utf-8') as f:
        raw = f.read()

    raw = deduplicate_sections(raw)
    raw = re.sub(r'（必達）', '', raw)   # 移除評分系統內部標記，不對外顯示

    _ticker = ticker or '7738'
    _company = company_name or f'（{_ticker}）'

    doc = Document()
    setup_document(doc)
    report_date = datetime.today().strftime("%Y-%m-%d")

    # ── Fetch live financial data ─────────────────────────────────────────────
    fin = None
    if HAS_FINANCIALS:
        try:
            print('  [cover] Fetching financial data…')
            fin = fetch_financials(_ticker, market)
            print(f'  [cover] 股價 {fin.fmt_price()}  市值 {fin.fmt_mktcap()}')
        except Exception as e:
            print(f'  [cover] Financial fetch failed: {e}')

    price_str  = fin.fmt_price()           if fin else 'N/A'
    mktcap_str = fin.fmt_mktcap()         if fin else 'N/A'
    sector_str = (fin.sector or '金融科技') if fin else '金融科技'
    market_str = (fin.market_type or '興櫃') if fin else '興櫃'
    rev_g_str  = fin.fmt_revenue_growth() if fin else 'N/A'
    pe_str     = fin.fmt_pe()             if fin else 'N/A'
    roe_str    = fin.fmt_roe()            if fin else 'N/A'
    peg_str    = fin.fmt_peg()            if fin else 'N/A'
    pb_str     = fin.fmt_pb()             if fin else 'N/A'
    eps_str    = fin.fmt_trailing_eps()   if fin else 'N/A'
    eps_q_str  = fin.fmt_eps_quarters()   if fin else 'N/A'
    w52_str    = fin.fmt_52w()            if fin else 'N/A'

    COVER_BG  = '0D1B3E'
    MID_BG    = '1E3D6E'
    LIGHT_BG  = 'E8F0F8'

    def _shd(cell_or_para, hex_color):
        """Apply background shading to a table cell."""
        tc = cell_or_para._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hex_color)
        tcPr.append(shd)

    def _cell_para(cell, text, size, bold=False,
                   color=(255,255,255), align=WD_ALIGN_PARAGRAPH.CENTER,
                   space_before=2, space_after=2):
        p = cell.paragraphs[0]
        p.clear()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after  = Pt(space_after)
        run = p.add_run(text)
        run.font.size  = Pt(size)
        run.font.bold  = bold
        run.font.color.rgb = RGBColor(*color)
        run.font.name  = '微軟正黑體'
        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微軟正黑體')
        return p

    def _gold_border_bottom(cell):
        """Add gold bottom border to a table cell."""
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcBdr = OxmlElement('w:tcBdr')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '12')
        bottom.set(qn('w:space'), '0')
        bottom.set(qn('w:color'), 'C9A84C')
        tcBdr.append(bottom)
        tcPr.append(tcBdr)

    # ── Cover header table (dark navy, full width) ────────────────────────────
    hdr_tbl = doc.add_table(rows=6, cols=1)
    hdr_tbl.style = 'Table Grid'
    rows_data = [
        (f'股票代號　{_ticker}　｜　' + market_str, 13, False, (0xA8, 0xD4, 0xF5)),
        (_company, 42, True,  (0xFF, 0xFF, 0xFF)),
        ('深度研究報告', 18, False, (0xA8, 0xD4, 0xF5)),
        (f'目前股價：{price_str}　｜　市值：{mktcap_str}', 11, False, (0xC0, 0xD8, 0xEE)),
        (f'產業：{sector_str}　｜　產出日期：{report_date}', 10, False, (0xC0, 0xD8, 0xEE)),
        ('TaiEquityautoresearch　｜　台灣市場深度研究平台', 8, False, (0x7B, 0xAA, 0xBF)),
    ]
    for i, (text, size, bold, color) in enumerate(rows_data):
        cell = hdr_tbl.rows[i].cells[0]
        _shd(cell, COVER_BG)
        sb = 14 if i == 0 else 2
        sa = 14 if i == len(rows_data)-1 else 2
        _cell_para(cell, text, size, bold=bold, color=color,
                   space_before=sb, space_after=sa)
        if i == 2:  # gold line under subtitle row
            _gold_border_bottom(cell)

    doc.add_paragraph()  # small gap

    # ── 5-column financial metrics table ──────────────────────────────────────
    met_tbl = doc.add_table(rows=3, cols=5)
    met_tbl.style = 'Table Grid'
    met_headers = ['TTM 營收成長', '本益比 P/E', 'ROE', '本益成長比 PEG', '本淨比 P/B']
    met_values  = [rev_g_str, pe_str, roe_str, peg_str, pb_str]
    met_labels  = ['YoY 年增率', 'Trailing', '股東報酬', 'PE ÷ EPS成長', '股價淨值']
    for ci in range(5):
        _shd(met_tbl.rows[0].cells[ci], MID_BG)
        _cell_para(met_tbl.rows[0].cells[ci], met_headers[ci], 8.5, bold=True,
                   color=(0xFF, 0xFF, 0xFF))
        _shd(met_tbl.rows[1].cells[ci], MID_BG)
        _cell_para(met_tbl.rows[1].cells[ci], met_values[ci], 14, bold=True,
                   color=(0xFF, 0xFF, 0xFF))
        _shd(met_tbl.rows[2].cells[ci], MID_BG)
        _cell_para(met_tbl.rows[2].cells[ci], met_labels[ci], 8,
                   color=(0x90, 0xBC, 0xD8))

    doc.add_paragraph()  # small gap

    # ── EPS / 52W 3-column table ──────────────────────────────────────────────
    eps_tbl = doc.add_table(rows=3, cols=3)
    eps_tbl.style = 'Table Grid'
    eps_headers = ['Trailing EPS (TTM)', '最近季度 EPS', '52 週股價區間']
    eps_values  = [eps_str, eps_q_str, f'NT$ {w52_str}']
    eps_labels  = ['近12個月每股盈餘', '興櫃：Q2 / Q4 申報', 'High – Low']
    for ci in range(3):
        _shd(eps_tbl.rows[0].cells[ci], LIGHT_BG)
        _cell_para(eps_tbl.rows[0].cells[ci], eps_headers[ci], 8.5, bold=True,
                   color=(0x1E, 0x3D, 0x6E))
        _shd(eps_tbl.rows[1].cells[ci], LIGHT_BG)
        _cell_para(eps_tbl.rows[1].cells[ci], eps_values[ci], 12, bold=True,
                   color=(0x0D, 0x1B, 0x3E))
        _shd(eps_tbl.rows[2].cells[ci], LIGHT_BG)
        _cell_para(eps_tbl.rows[2].cells[ci], eps_labels[ci], 8,
                   color=(0x55, 0x80, 0xA0))

    doc.add_paragraph()  # small gap

    # ── Disclaimer row ────────────────────────────────────────────────────────
    disc_tbl = doc.add_table(rows=1, cols=1)
    disc_tbl.style = 'Table Grid'
    _shd(disc_tbl.rows[0].cells[0], COVER_BG)
    _cell_para(disc_tbl.rows[0].cells[0],
               '本報告由 TaiEquityautoresearch 自動生成，僅供參考，不構成投資建議。財務數據來源：Yahoo Finance。',
               7.5, color=(0x7B, 0xAA, 0xBF), space_before=4, space_after=4)

    # Page break after cover
    doc.add_page_break()

    lines = raw.split('\n')
    i = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Blank
        if not stripped:
            i += 1
            continue

        # Separator → horizontal rule (just a styled paragraph)
        if is_separator(line):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            pPr = p._p.get_or_add_pPr()
            pb = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '4')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), 'CCCCCC')
            pb.append(bottom)
            pPr.append(pb)
            i += 1
            continue

        # Table
        if is_table_row(line):
            table_lines = []
            while i < len(lines) and is_table_row(lines[i]):
                table_lines.append(lines[i])
                i += 1
            rows = parse_table_rows(table_lines)
            if not rows:
                continue
            col_count = max(len(r) for r in rows)
            rows = [r + [''] * (col_count - len(r)) for r in rows]

            tbl = doc.add_table(rows=len(rows), cols=col_count)
            tbl.style = 'Table Grid'

            for ri, row in enumerate(rows):
                for ci, cell_text in enumerate(row):
                    cell = tbl.cell(ri, ci)
                    cell.paragraphs[0].clear()
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after  = Pt(2)
                    if ri == 0:
                        # Header row: dark blue bg, white bold text
                        set_cell_bg(cell, '0F3460')
                        run = p.add_run(re.sub(r'\*\*(.+?)\*\*', r'\1', cell_text))
                        run.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        run.font.name = '微軟正黑體'
                        run._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')
                    else:
                        if ri % 2 == 0:
                            set_cell_bg(cell, 'F0F4F8')
                        apply_inline_to_para(p, cell_text)
                        for run in p.runs:
                            run.font.name = '微軟正黑體'
                            run._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')

            doc.add_paragraph()  # spacer after table
            continue

        # Blockquote
        if stripped.startswith('> '):
            quote_text = stripped[2:]
            p = doc.add_paragraph(style='Normal')
            p.paragraph_format.left_indent  = Cm(1)
            p.paragraph_format.right_indent = Cm(0.5)
            p.paragraph_format.space_after  = Pt(4)
            # Grey left border
            pPr = p._p.get_or_add_pPr()
            pb = OxmlElement('w:pBdr')
            left = OxmlElement('w:left')
            left.set(qn('w:val'), 'single')
            left.set(qn('w:sz'), '12')
            left.set(qn('w:space'), '4')
            left.set(qn('w:color'), '888888')
            pb.append(left)
            pPr.append(pb)
            apply_inline_to_para(p, quote_text)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
                run.font.name = '微軟正黑體'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')
            i += 1
            continue

        # H1
        if re.match(r'^# ', line):
            text = line[2:].strip()
            p = doc.add_heading(level=1)
            p.clear()
            apply_inline_to_para(p, text)
            for run in p.runs:
                run.font.name = '微軟正黑體'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')
                run.font.size = Pt(15)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x16, 0x21, 0x3e)
            i += 1
            continue

        # H2
        if re.match(r'^## ', line):
            text = line[3:].strip()
            p = doc.add_heading(level=2)
            p.clear()
            apply_inline_to_para(p, text)
            for run in p.runs:
                run.font.name = '微軟正黑體'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')
                run.font.size = Pt(13)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x0f, 0x34, 0x60)
            i += 1
            continue

        # H3
        if re.match(r'^### ', line):
            text = line[4:].strip()
            p = doc.add_heading(level=3)
            p.clear()
            apply_inline_to_para(p, text)
            for run in p.runs:
                run.font.name = '微軟正黑體'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')
                run.font.size = Pt(11)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x53, 0x34, 0x83)
            i += 1
            continue

        # H4
        if re.match(r'^#### ', line):
            text = line[5:].strip()
            p = doc.add_paragraph(style='Normal')
            p.paragraph_format.space_before = Pt(6)
            apply_inline_to_para(p, text, is_bold_default=True)
            for run in p.runs:
                run.font.name = '微軟正黑體'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')
            i += 1
            continue

        # Italic / emphasis line (e.g. *注：...*)
        if stripped.startswith('*') and stripped.endswith('*') and not stripped.startswith('**'):
            inner = stripped[1:-1]
            p = doc.add_paragraph(style='Normal')
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(inner)
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            run.font.name = '微軟正黑體'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')
            i += 1
            continue

        # List item
        if re.match(r'^[\-\*] ', stripped) or re.match(r'^\d+\. ', stripped):
            bullet_text = re.sub(r'^[\-\*\d\.]+\s*', '', stripped)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            apply_inline_to_para(p, bullet_text)
            for run in p.runs:
                run.font.name = '微軟正黑體'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')
            i += 1
            continue

        # Normal paragraph
        p = doc.add_paragraph(style='Normal')
        p.paragraph_format.space_after = Pt(4)
        apply_inline_to_para(p, stripped)
        for run in p.runs:
            run.font.name = '微軟正黑體'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), '微軟正黑體')
        i += 1

    doc.save(out_path)
    print(f'[OK] Word generated: {out_path}')
    size_kb = os.path.getsize(out_path) // 1024
    print(f'  Size: {size_kb} KB')


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('--ticker', default='7738')
    parser.add_argument('--company', default=None)
    parser.add_argument('--market', default='TWO')
    args = parser.parse_args()
    base = rf'C:\Users\機動小隊\TaiEquityautoresearch\data\companies\{args.ticker}'
    md  = rf'{base}\{args.ticker}_Initial_MAX.md'
    out = rf'{base}\{args.ticker}_Report.docx'
    build_word(md, out, ticker=args.ticker, company_name=args.company, market=args.market)
