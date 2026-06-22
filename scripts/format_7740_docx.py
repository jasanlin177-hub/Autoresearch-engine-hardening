from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(r"C:\Users\機動小隊\TaiEquityautoresearch")
DOCX_PATH = ROOT / "data" / "companies" / "7740" / "7740_Initial_MAX.docx"
OUTPUT_PATH = ROOT / "data" / "companies" / "7740" / "7740_Initial_MAX_broker.docx"


def set_run_font(run, east_asia: str, latin: str, size: int, bold: bool | None = None):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def insert_paragraph_before(paragraph, text="", style=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p = new_p
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def add_toc(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    hint = OxmlElement("w:t")
    hint.text = "更新目錄請在 Word 中按右鍵並選擇「更新欄位」。"
    fld_sep.append(hint)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def apply_styles(doc: Document):
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微軟正黑體")
    normal.font.size = Pt(11)

    for style_name, size in [("Title", 24), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微軟正黑體")
        style.font.size = Pt(size)
        style.font.bold = True


def add_cover_and_toc(doc: Document):
    first = doc.paragraphs[0]
    for _ in range(8):
        insert_paragraph_before(first, "")

    p = insert_paragraph_before(first, "熙特爾新能源（7740）正式投資備忘錄", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        set_run_font(run, "微軟正黑體", "Times New Roman", 24, True)

    p2 = insert_paragraph_before(first, "TaiEquityautoresearch Equity Research", style="Subtitle")
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p2.runs:
        set_run_font(run, "微軟正黑體", "Times New Roman", 13)

    p2b = insert_paragraph_before(first, "台灣儲能系統整合商深度研究")
    p2b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p2b.runs:
        set_run_font(run, "微軟正黑體", "Times New Roman", 12)

    p3 = insert_paragraph_before(first, "公司：熙特爾新能源股份有限公司")
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4 = insert_paragraph_before(first, "股票代號：7740")
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p5 = insert_paragraph_before(first, "報告日期：2026-05-17")
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p6 = insert_paragraph_before(first, "分析框架：基本面 / 風險附錄 / 投資委員會摘要")
    p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p7 = insert_paragraph_before(first, "作者：Codex x TaiEquityautoresearch")
    p7.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p8 = insert_paragraph_before(first, "用途：內部研究討論文件")
    p8.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for para in [p3, p4, p5, p6, p7, p8]:
        for run in para.runs:
            set_run_font(run, "微軟正黑體", "Times New Roman", 11)

    page_break = insert_paragraph_before(first, "")
    page_break.add_run().add_break()

    disclaimer_title = insert_paragraph_before(first, "免責聲明", style="Heading 1")
    for run in disclaimer_title.runs:
        set_run_font(run, "微軟正黑體", "Times New Roman", 16, True)

    disclaimer_lines = [
        "本文件僅供研究、教育與內部討論使用，不構成任何買賣證券之要約、邀約或投資建議。",
        "本文件內容主要整理自公司官方年報、法說簡報、公開說明書、股東會資料及公開市場資訊，雖力求正確，仍可能因揭露時點、資料更新或轉述整理而產生差異。",
        "所有估值、情境分析、投資判斷與風險評估均屬研究觀點，不保證未來報酬或價格表現，使用者應自行判斷並承擔投資風險。",
    ]
    for line in disclaimer_lines:
        para = insert_paragraph_before(first, line)
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in para.runs:
            set_run_font(run, "微軟正黑體", "Times New Roman", 10)

    disclaimer_break = insert_paragraph_before(first, "")
    disclaimer_break.add_run().add_break()

    toc_title = insert_paragraph_before(first, "目錄", style="Heading 1")
    for run in toc_title.runs:
        set_run_font(run, "微軟正黑體", "Times New Roman", 16, True)

    toc_paragraph = insert_paragraph_before(first, "")
    add_toc(toc_paragraph)

    toc_break = insert_paragraph_before(first, "")
    toc_break.add_run().add_break()


def add_header_footer(doc: Document):
    for section in doc.sections:
        header = section.header
        header_para = header.paragraphs[0]
        header_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        header_para.text = "TaiEquityautoresearch | 熙特爾新能源（7740）"
        for run in header_para.runs:
            set_run_font(run, "微軟正黑體", "Times New Roman", 9)

        footer = section.footer
        footer_para = footer.paragraphs[0]
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_para.text = "第 "
        add_page_number(footer_para)
        footer_para.add_run(" 頁")
        for run in footer_para.runs:
            set_run_font(run, "微軟正黑體", "Times New Roman", 9)


def clean_spacing(doc: Document):
    for paragraph in doc.paragraphs:
        fmt = paragraph.paragraph_format
        fmt.space_after = Pt(6)
        fmt.line_spacing = 1.25
        if paragraph.style.name.startswith("Heading"):
            fmt.space_before = Pt(12)
            fmt.space_after = Pt(5)
        for run in paragraph.runs:
            if paragraph.style.name == "Title":
                set_run_font(run, "微軟正黑體", "Times New Roman", 24, True)
            elif paragraph.style.name.startswith("Heading 1"):
                set_run_font(run, "微軟正黑體", "Times New Roman", 16, True)
            elif paragraph.style.name.startswith("Heading 2"):
                set_run_font(run, "微軟正黑體", "Times New Roman", 14, True)
            elif paragraph.style.name.startswith("Heading 3"):
                set_run_font(run, "微軟正黑體", "Times New Roman", 12, True)
            else:
                set_run_font(run, "微軟正黑體", "Times New Roman", 11)


def style_tables(doc: Document):
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(2)
                    paragraph.paragraph_format.line_spacing = 1.1
                    for run in paragraph.runs:
                        set_run_font(run, "微軟正黑體", "Times New Roman", 10)


def main():
    doc = Document(DOCX_PATH)
    apply_styles(doc)
    add_cover_and_toc(doc)
    add_header_footer(doc)
    clean_spacing(doc)
    style_tables(doc)
    doc.save(OUTPUT_PATH)


if __name__ == "__main__":
    main()
